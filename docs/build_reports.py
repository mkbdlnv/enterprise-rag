"""Builds the assignment report: Markdown -> DOCX (pandoc) -> required formatting.

    python docs/build_reports.py      # needs pandoc and python-docx

The assignment asks for 12 pt text with 1.5 line spacing; pandoc's default
template uses other fonts and sizes, so every style is normalised here.
"""

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parent
FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)


def set_font(style, size=None, bold=None):
    font = style.font
    font.name = FONT
    font.color.rgb = BLACK
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    # Theme fonts would override the explicit name in Word.
    fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
    if fonts is not None:
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), FONT)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            fonts.attrib.pop(qn(attr), None)


def table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tbl_pr.append(borders)


def format_docx(path: Path):
    doc = Document(path)
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.top_margin = section.bottom_margin = Cm(2)

    sizes = {"Title": (16, True), "Heading 1": (14, True), "Heading 2": (12, True)}
    for style in doc.styles:
        if style.type != 1:  # paragraph styles only
            continue
        size, bold = sizes.get(style.name, (12, None))
        set_font(style, size, bold)
        style.font.italic = False if style.name in sizes else style.font.italic
        pf = style.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(6)
        if style.name in ("Title", "Subtitle", "Author", "Date"):
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if style.name in ("Body Text", "First Paragraph", "Compact"):
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if style.name.startswith("Heading"):
            pf.space_before = Pt(12)
            pf.keep_with_next = True

    in_refs = False
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            in_refs = p.text.strip() in ("References", "Список литературы")
        elif in_refs:
            # The reference list is set single-spaced, as usual for bibliographies.
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            run.font.name = FONT
            run.font.color.rgb = BLACK
            if p.style.name not in ("Title",) and not p.style.name.startswith("Heading"):
                run.font.size = Pt(12)
        if p.style.name in ("Image Caption", "Captioned Figure"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(11)
                run.font.italic = True

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(table)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.15
                    for run in p.runs:
                        run.font.name = FONT
                        run.font.size = Pt(11)
    doc.save(path)


def main():
    for name in ("report",):
        out = DOCS / f"{name}.docx"
        subprocess.run(
            ["pandoc", f"{name}.md", "-o", out.name, "--resource-path", "."],
            cwd=DOCS,
            check=True,
        )
        format_docx(out)
        print("built", out)


if __name__ == "__main__":
    main()
